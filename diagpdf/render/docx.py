"""Word output, including embedded fonts and running headers."""

from __future__ import annotations

import re
import uuid
import zipfile
from io import BytesIO
from pathlib import Path
from xml.sax.saxutils import escape

from ..chars import _ALPHA_REQUIRED_CHARS, _MERIDA_REQUIRED_CHARS
from ..fen import fen_to_diagram, parse_fen
from ..fonts import (
    _board_font_family_name,
    check_embedding_permitted,
    font_embedding_allowed,
    resolve_board_font,
)
from ..i18n import _answers_heading, _lichess_label
from ..layouts import LAYOUTS
from ..log import _warn
from ..model import PositionDict
from ..options import RenderOptions
from ..parsers import _clean_moves
from ..resources import _resource
from .common import (
    _anchor_id,
    _chapter_groups,
    _figurine_family_name,
    _figurine_font_path,
    _figurine_segments,
    _format_number,
    _has_embedded_side_indicator,
    _layout_columns,
    _lichess_analysis_url,
    _make_title,
    _numbering,
    _side_indicator_text,
    _split_title_number_prefix,
)


def _docx_add_bookmark(paragraph, name: str, bookmark_id: int) -> None:
    """Add an internal bookmark anchor to a python-docx paragraph."""
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    start = OxmlElement('w:bookmarkStart')
    start.set(qn('w:id'), str(bookmark_id))
    start.set(qn('w:name'), name)

    end = OxmlElement('w:bookmarkEnd')
    end.set(qn('w:id'), str(bookmark_id))

    paragraph._p.append(start)
    paragraph._p.append(end)


def _docx_set_run_font(run, font_name: str, font_size_pt: float | None = None) -> None:
    """Force a run to use the given font across all Word font slots."""
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Pt

    run.font.name = font_name
    if font_size_pt is not None:
        run.font.size = Pt(font_size_pt)

    r_pr = run._element.get_or_add_rPr()
    r_fonts = r_pr.rFonts
    if r_fonts is None:
        r_fonts = OxmlElement('w:rFonts')
        r_pr.insert(0, r_fonts)
    for key in ('ascii', 'hAnsi', 'eastAsia', 'cs'):
        r_fonts.set(qn(f'w:{key}'), font_name)


def _docx_add_board_paragraph(
    paragraph,
    diag_lines: list[str],
    font_name: str,
    font_size_pt: float,
) -> None:
    """Render the diagram as text lines in the chess font."""
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
    from docx.shared import Pt

    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    paragraph.paragraph_format.line_spacing = 1.0
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run('\n'.join(diag_lines or ['']))
    _docx_set_run_font(run, font_name, font_size_pt)


def _docx_font_embed_policy(font_path: Path) -> tuple[bool, bool]:
    """Deprecated alias kept for compatibility — see font_embedding_allowed."""
    return font_embedding_allowed(font_path)


def _docx_prepare_font_payload(font_path: Path, used_chars: set[str]) -> tuple[bytes, bool]:
    """Return (font_bytes, subsetted) for DOCX embedding.

    The caller is responsible for checking that embedding is permitted; this
    function only decides whether the payload can be subsetted.
    """
    _, can_subset = font_embedding_allowed(font_path)

    if can_subset and used_chars:
        try:
            from fontTools import subset
            from fontTools.ttLib import TTFont

            options = subset.Options()
            options.desubroutinize = False
            options.notdef_outline = True
            options.recommended_glyphs = True

            font = TTFont(str(font_path))
            subsetter = subset.Subsetter(options=options)
            subsetter.populate(text=''.join(sorted(used_chars)))
            subsetter.subset(font)
            buf = BytesIO()
            font.save(buf)
            font.close()
            return buf.getvalue(), True
        except Exception:
            pass

    return font_path.read_bytes(), False


def _docx_font_metadata_xml(font_path: Path) -> str:
    """Build Word font metadata tags for fontTable.xml."""
    try:
        from fontTools.ttLib import TTFont

        font = TTFont(str(font_path))
        os2 = font['OS/2'] if 'OS/2' in font else None
        post = font['post'] if 'post' in font else None
        cmap_codepoints: set[int] = set()
        if 'cmap' in font:
            for table in font['cmap'].tables:
                cmap_codepoints.update(table.cmap.keys())
        is_symbol_like = (
            any(0xF000 + ord(ch) in cmap_codepoints for ch in _MERIDA_REQUIRED_CHARS | _ALPHA_REQUIRED_CHARS)
            or (os2 is not None and all(int(getattr(os2, name, 0)) == 0 for name in ('ulUnicodeRange1', 'ulUnicodeRange2', 'ulUnicodeRange3', 'ulUnicodeRange4')))
        )

        family_map = {
            1: 'roman',
            2: 'swiss',
            3: 'modern',
            4: 'script',
            5: 'decorative',
        }

        parts: list[str] = []
        if os2 is not None and hasattr(os2, 'panose'):
            panose = os2.panose
            vals = [
                getattr(panose, attr, 0)
                for attr in (
                    'bFamilyType', 'bSerifStyle', 'bWeight', 'bProportion', 'bContrast',
                    'bStrokeVariation', 'bArmStyle', 'bLetterForm', 'bMidline', 'bXHeight',
                )
            ]
            parts.append(f'<w:panose1 w:val="{"".join(f"{v:02X}" for v in vals)}"/>')

        charset = '02' if is_symbol_like else '00'
        parts.append(f'<w:charset w:val="{charset}"/>')

        family = 'auto'
        if os2 is not None:
            family = family_map.get((int(getattr(os2, 'sFamilyClass', 0)) >> 8) & 0xFF, 'auto')
        if family == 'auto' and is_symbol_like:
            family = 'swiss'
        parts.append(f'<w:family w:val="{family}"/>')

        pitch = 'fixed' if post is not None and int(getattr(post, 'isFixedPitch', 0)) else 'variable'
        parts.append(f'<w:pitch w:val="{pitch}"/>')

        if is_symbol_like:
            parts.append(
                '<w:sig '
                'w:usb0="00000000" w:usb1="10000000" w:usb2="00000000" w:usb3="00000000" '
                'w:csb0="80000000" w:csb1="00000000"/>'
            )
        elif os2 is not None:
            usb = [int(getattr(os2, name, 0)) for name in ('ulUnicodeRange1', 'ulUnicodeRange2', 'ulUnicodeRange3', 'ulUnicodeRange4')]
            if hasattr(os2, 'ulCodePageRange1'):
                csb = [int(getattr(os2, name, 0)) for name in ('ulCodePageRange1', 'ulCodePageRange2')]
            else:
                csb = [0, 0]
            parts.append(
                '<w:sig '
                f'w:usb0="{usb[0]:08X}" w:usb1="{usb[1]:08X}" w:usb2="{usb[2]:08X}" w:usb3="{usb[3]:08X}" '
                f'w:csb0="{csb[0]:08X}" w:csb1="{csb[1]:08X}"/>'
            )

        font.close()
        return ''.join(parts)
    except Exception:
        return '<w:charset w:val="00"/><w:family w:val="auto"/><w:pitch w:val="variable"/>'


def _docx_obfuscate_font_bytes(font_bytes: bytes, font_key: str) -> bytes:
    """Obfuscate a font payload for the DOCX .odttf part."""
    hex_key = font_key.strip('{}').replace('-', '')
    key = bytes.fromhex(''.join(reversed([hex_key[i:i + 2] for i in range(0, len(hex_key), 2)])))
    data = bytearray(font_bytes)
    for idx in range(min(32, len(data))):
        data[idx] ^= key[idx % len(key)]
    return bytes(data)


def _docx_embed_font(docx_path: Path, font_name: str, font_path: Path, used_chars: set[str]) -> None:
    """Embed the chess font inside a DOCX package."""
    from xml.sax.saxutils import escape

    font_bytes, subsetted = _docx_prepare_font_payload(font_path, used_chars)
    font_key = '{' + str(uuid.uuid4()).upper() + '}'
    obfuscated = _docx_obfuscate_font_bytes(font_bytes, font_key)
    metadata_xml = _docx_font_metadata_xml(font_path)

    with zipfile.ZipFile(docx_path, 'r') as zin:
        files = {name: zin.read(name) for name in zin.namelist()}

    font_table_bytes = files.get('word/fontTable.xml')
    settings_bytes = files.get('word/settings.xml')
    content_types_bytes = files.get('[Content_Types].xml')
    if not font_table_bytes or not settings_bytes or not content_types_bytes:
        raise RuntimeError('DOCX package is missing required XML parts for font embedding.')

    font_table_xml = font_table_bytes.decode('utf-8')
    settings_xml = settings_bytes.decode('utf-8')
    content_types_xml = content_types_bytes.decode('utf-8')
    rels_path = 'word/_rels/fontTable.xml.rels'
    rels_xml = files.get(
        rels_path,
        b'<?xml version="1.0" encoding="UTF-8" standalone="yes"?><Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"/>',
    ).decode('utf-8')

    used_ids = set(re.findall(r'Id="(rId\d+)"', rels_xml))
    next_id = 1
    while f'rId{next_id}' in used_ids:
        next_id += 1
    rel_id = f'rId{next_id}'

    used_targets = set(re.findall(r'Target="([^"]+)"', rels_xml))
    next_font_idx = 1
    while f'fonts/font{next_font_idx}.odttf' in used_targets:
        next_font_idx += 1
    font_target = f'fonts/font{next_font_idx}.odttf'

    font_attr = escape(font_name, {'"': '&quot;'})
    embed_tag = f'<w:embedRegular r:id="{rel_id}" w:fontKey="{font_key}"/>'
    full_font_inner = metadata_xml + embed_tag

    block_pat = re.compile(
        rf'(<w:font\b[^>]*w:name="{re.escape(font_attr)}"[^>]*>)(.*?)(</w:font>)',
        re.DOTALL,
    )
    self_pat = re.compile(rf'<w:font\b([^>]*?)w:name="{re.escape(font_attr)}"([^>]*)/>')

    if block_pat.search(font_table_xml):
        def _replace_font_block(match: re.Match[str]) -> str:
            inner = re.sub(r'<w:embedRegular\b[^>]*/>', '', match.group(2))
            if '<w:panose1' not in inner and '<w:charset' not in inner and '<w:family' not in inner:
                inner = metadata_xml + inner
            return match.group(1) + inner + embed_tag + match.group(3)
        font_table_xml = block_pat.sub(_replace_font_block, font_table_xml, count=1)
    elif self_pat.search(font_table_xml):
        font_table_xml = self_pat.sub(
            rf'<w:font\1w:name="{font_attr}"\2>{full_font_inner}</w:font>',
            font_table_xml,
            count=1,
        )
    else:
        font_table_xml = font_table_xml.replace(
            '</w:fonts>',
            f'<w:font w:name="{font_attr}">{full_font_inner}</w:font></w:fonts>',
            1,
        )

    rel_tag = (
        f'<Relationship Id="{rel_id}" '
        'Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/font" '
        f'Target="{font_target}"/>'
    )
    if re.search(r'<Relationships\b[^>]*/>\s*$', rels_xml):
        rels_xml = re.sub(
            r'<Relationships\b([^>]*)/>\s*$',
            rf'<Relationships\1>{rel_tag}</Relationships>',
            rels_xml,
            count=1,
        )
    elif '</Relationships>' in rels_xml:
        rels_xml = rels_xml.replace('</Relationships>', rel_tag + '</Relationships>', 1)
    else:
        rels_xml = (
            '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
            '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
            f'{rel_tag}</Relationships>'
        )

    if '<w:embedTrueTypeFonts' not in settings_xml:
        settings_xml = settings_xml.replace('</w:settings>', '<w:embedTrueTypeFonts/></w:settings>', 1)
    if subsetted and '<w:saveSubsetFonts' not in settings_xml:
        settings_xml = settings_xml.replace('</w:settings>', '<w:saveSubsetFonts/></w:settings>', 1)

    if 'Extension="odttf"' not in content_types_xml:
        content_types_xml = content_types_xml.replace(
            '</Types>',
            '<Default Extension="odttf" ContentType="application/vnd.openxmlformats-officedocument.obfuscatedFont"/></Types>',
            1,
        )

    files['word/fontTable.xml'] = font_table_xml.encode('utf-8')
    files[rels_path] = rels_xml.encode('utf-8')
    files['word/settings.xml'] = settings_xml.encode('utf-8')
    files['[Content_Types].xml'] = content_types_xml.encode('utf-8')
    files[f'word/{font_target}'] = obfuscated

    with zipfile.ZipFile(docx_path, 'w', compression=zipfile.ZIP_DEFLATED) as zout:
        for name, data in files.items():
            zout.writestr(name, data)

def _docx_add_hyperlink(
    paragraph,
    text: str,
    url: str = '',
    anchor: str = '',
    bold: bool = False,
    font_size_pt: float | None = None,
) -> None:
    """Add a clickable hyperlink to a python-docx paragraph."""
    from docx.opc.constants import RELATIONSHIP_TYPE as RT
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    hyperlink = OxmlElement('w:hyperlink')
    if anchor:
        hyperlink.set(qn('w:anchor'), anchor)
        hyperlink.set(qn('w:history'), '1')
    else:
        part = paragraph.part
        r_id = part.relate_to(url, RT.HYPERLINK, is_external=True)
        hyperlink.set(qn('r:id'), r_id)

    new_run = OxmlElement('w:r')
    r_pr = OxmlElement('w:rPr')

    color = OxmlElement('w:color')
    color.set(qn('w:val'), '0000FF')
    r_pr.append(color)

    underline = OxmlElement('w:u')
    underline.set(qn('w:val'), 'single')
    r_pr.append(underline)
    if bold:
        bold_el = OxmlElement('w:b')
        r_pr.append(bold_el)
    if font_size_pt is not None:
        sz = OxmlElement('w:sz')
        sz.set(qn('w:val'), str(int(round(font_size_pt * 2))))
        r_pr.append(sz)

    text_el = OxmlElement('w:t')
    text_el.text = text

    new_run.append(r_pr)
    new_run.append(text_el)
    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)

def _docx_add_field(paragraph, instruction: str) -> None:
    """Insert a Word field such as PAGE or NUMPAGES into a paragraph."""
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    run = OxmlElement('w:r')
    begin = OxmlElement('w:fldChar')
    begin.set(qn('w:fldCharType'), 'begin')
    instr = OxmlElement('w:instrText')
    instr.set(qn('xml:space'), 'preserve')
    instr.text = f' {instruction} '
    end = OxmlElement('w:fldChar')
    end.set(qn('w:fldCharType'), 'end')
    run.append(begin)
    run.append(instr)
    run.append(end)
    paragraph._p.append(run)


def _docx_write_running_text(paragraph, template: str, chapter: str) -> None:
    """Write header/footer text, turning {page}/{total} into live Word fields."""
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    text = template.replace('{chapter}', chapter)
    for piece in re.split(r'(\{page\}|\{total\})', text):
        if piece == '{page}':
            _docx_add_field(paragraph, 'PAGE')
        elif piece == '{total}':
            _docx_add_field(paragraph, 'NUMPAGES')
        elif piece:
            paragraph.add_run(piece)


# ── Watermark ────────────────────────────────────────────────────────────────
# A watermark in Word is not text: it is a WordArt drawing anchored in the
# header, which is the only part of the document that repeats on every page.
# The shape type below is the one Word itself writes -- the formulas are what
# bend the text along the path, and Word rejects the shape without them.
_VML_DECLS = (
    'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
    'xmlns:v="urn:schemas-microsoft-com:vml" '
    'xmlns:o="urn:schemas-microsoft-com:office:office"'
)

_WORDART_SHAPETYPE = (
    '<v:shapetype id="_x0000_t136" coordsize="21600,21600" o:spt="136" adj="10800" '
    'path="m@7,0l@8,0m@5,21600l@6,21600e">'
    '<v:formulas>'
    '<v:f eqn="sum #0 0 10800"/><v:f eqn="prod #0 2 1"/><v:f eqn="sum 21600 0 @1"/>'
    '<v:f eqn="sum 0 0 @2"/><v:f eqn="sum 21600 0 @3"/><v:f eqn="if @0 @3 0"/>'
    '<v:f eqn="if @0 21600 @1"/><v:f eqn="if @0 0 @2"/><v:f eqn="if @0 @4 21600"/>'
    '<v:f eqn="mid @5 @6"/><v:f eqn="mid @8 @5"/><v:f eqn="mid @7 @8"/>'
    '<v:f eqn="mid @6 @7"/><v:f eqn="sum @6 0 @5"/>'
    '</v:formulas>'
    '<v:path textpathok="t" o:connecttype="custom" '
    'o:connectlocs="@9,0;@10,10800;@11,21600;@12,10800" o:connectangles="270,180,90,0"/>'
    '<v:textpath on="t" fitshape="t"/>'
    '<v:handles><v:h position="#0,bottomRight" xrange="6629,14971"/></v:handles>'
    '<o:lock v:ext="edit" text="t" shapetype="t"/>'
    '</v:shapetype>'
)


def _docx_add_watermark(section, text: str, opacity: int, page) -> None:
    """Stamp *text* diagonally across every page of *section*.

    The shape goes into the header paragraph that is already there rather than
    into one of its own: it is positioned absolutely, so it costs no line, and a
    second header paragraph would push the body down by an empty line.

    Grey at the requested opacity, to match what the PDF stamps -- Word's own
    "washout" is a fixed colour that no option could then move.
    """
    from docx.oxml import parse_xml

    text = text.strip()
    if not text:
        return
    # Wide enough to cross the sheet once turned: the shape is the box the text
    # is fitted into, and 315 degrees is Word's own diagonal.
    width_pt = page.width * 2.8346 * 0.85
    height_pt = width_pt / 4.5
    quoted = escape(text, {'"': '&quot;'})
    shape = (
        f'<w:r {_VML_DECLS}><w:pict>'
        + _WORDART_SHAPETYPE
        + '<v:shape id="DiagPDFWatermark" o:spid="_x0000_s2049" type="#_x0000_t136" '
        f'style="position:absolute;margin-left:0;margin-top:0;'
        f'width:{width_pt:.2f}pt;height:{height_pt:.2f}pt;rotation:315;z-index:-251654144;'
        'mso-position-horizontal:center;mso-position-horizontal-relative:margin;'
        'mso-position-vertical:center;mso-position-vertical-relative:margin" '
        'o:allowincell="f" fillcolor="#808080" stroked="f">'
        f'<v:fill opacity="{opacity / 100:.4f}"/>'
        f'<v:textpath style="font-family:&quot;Calibri&quot;;font-size:1pt" string="{quoted}"/>'
        '</v:shape></w:pict></w:r>'
    )
    section.header.paragraphs[0]._p.append(parse_xml(shape))


def _docx_set_columns(section, count: int) -> None:
    """Lay a section out in *count* newspaper columns."""
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    sect_pr = section._sectPr
    cols = sect_pr.find(qn('w:cols'))
    if cols is None:
        cols = OxmlElement('w:cols')
        sect_pr.append(cols)
    cols.set(qn('w:num'), str(max(1, count)))
    cols.set(qn('w:space'), '425')


def _docx_notation_paragraph(container, index: int, numbered: bool, side: str, width_twips: int):
    """Add one blank notation line, drawn with underscore tab leaders."""
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Pt

    paragraph = container.add_paragraph()
    paragraph.style = container.part.document.styles['diagpdf-notation']
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(2)

    # Black to move: the first move number belongs to White, so its blank is
    # left out and only Black's is offered.
    skip_first = numbered and index == 0 and side == 'b'
    stops = [width_twips // 2, width_twips] if numbered else [width_twips]
    leaders = ['none' if skip_first else 'underscore', 'underscore'] if numbered else ['underscore']

    p_pr = paragraph._p.get_or_add_pPr()
    tabs = OxmlElement('w:tabs')
    for position, leader in zip(stops, leaders, strict=True):
        tab = OxmlElement('w:tab')
        tab.set(qn('w:val'), 'left')
        tab.set(qn('w:leader'), leader)
        tab.set(qn('w:pos'), str(position))
        tabs.append(tab)
    p_pr.append(tabs)

    if numbered:
        run = paragraph.add_run(f'{index + 1}.')
        run.font.size = Pt(8)
    paragraph.add_run('\t' * len(stops))
    return paragraph


def _docx_add_moves(paragraph, moves: str, opts: RenderOptions, figurine_family: str,
                    used_figurine_chars: set[str]) -> None:
    """Add movetext, switching piece letters to the figurine font."""
    if not figurine_family:
        paragraph.add_run(moves)
        return
    for font, run_text in _figurine_segments(moves, '', 'figurine'):
        run = paragraph.add_run(run_text)
        if font == 'figurine':
            _docx_set_run_font(run, figurine_family)
            used_figurine_chars.update(run_text)


def _docx_ensure_notation_style(doc) -> None:
    """Register the paragraph style used by the blank notation lines."""
    from docx.enum.style import WD_STYLE_TYPE
    if 'diagpdf-notation' in [s.name for s in doc.styles]:
        return
    style = doc.styles.add_style('diagpdf-notation', WD_STYLE_TYPE.PARAGRAPH)
    style.base_style = doc.styles['Normal']
    style.hidden = False
    style.quick_style = False

def generate_docx(positions: list[PositionDict], opts: RenderOptions | dict,
                  out_path, progress=None) -> None:
    """Render a DOCX document with diagram text, links, and optional answers."""
    try:
        from docx import Document
        from docx.enum.section import WD_ORIENT
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.shared import Pt
    except ImportError as e:
        raise RuntimeError('python-docx required — install with: pip install python-docx') from e

    opts = RenderOptions.coerce(opts)
    valid = [p for p in positions if p.get('fen', '').strip()]
    total = len(valid)
    font_name = opts.font
    font_family_name = _board_font_family_name(font_name)
    title_template = opts.title_template
    title_mode = opts.title_mode
    title_custom = opts.title_custom
    lichess_link = opts.lichess_link
    answers_section = opts.answers_section
    answers_title = _answers_heading(opts)
    coords = opts.coords
    flip = opts.flip
    flip_auto = opts.flip_auto
    symbol = opts.symbol
    border_style = opts.border_style
    board_font_size_pt = opts.font_size if opts.font_size > 0 else LAYOUTS[opts.layout_idx][6]
    fonts_dir = _resource('Fonts')
    font_file = resolve_board_font(font_name)
    docx_font_path = fonts_dir / font_file

    lines_count = opts.lines_count
    lines_numbered = opts.lines_mode == 'numbered'
    answers_cols = opts.answers_cols
    columns = _layout_columns(opts)
    figurine_family = _figurine_family_name(opts) if answers_section else ''

    doc = Document()
    section = doc.sections[0]
    from docx.shared import Mm
    page = opts.page
    section.page_width = Mm(page.width)
    section.page_height = Mm(page.height)
    section.orientation = WD_ORIENT.LANDSCAPE if page.landscape else WD_ORIENT.PORTRAIT
    section.top_margin = Mm(page.margin_tb)
    section.bottom_margin = Mm(page.margin_tb)
    section.left_margin = Mm(page.margin_lr)
    section.right_margin = Mm(page.margin_lr)

    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Calibri'
    normal_style.font.size = Pt(11)
    _docx_ensure_notation_style(doc)

    # Running header and footer, with live page-number fields.
    header_text = opts.header_text
    footer_text = opts.footer_text
    first_chapter = next((p.get('chapter', '') for p in valid if p.get('chapter')), '')
    if header_text:
        _docx_write_running_text(section.header.paragraphs[0], header_text, first_chapter)
    if footer_text:
        _docx_write_running_text(section.footer.paragraphs[0], footer_text, first_chapter)
    _docx_add_watermark(section, opts.watermark, opts.watermark_opacity, page)

    usable_twips = int(page.usable_width * 56.6929)   # sheet width minus the margins
    cell_twips = max(720, usable_twips // max(1, columns))

    numbers = _numbering(valid, opts)
    answers: list[tuple[int, str, str, str]] = []
    used_board_chars: set[str] = set()
    used_figurine_chars: set[str] = set()
    bookmark_id = 1
    rendered = 0

    def render_position(container, pos_idx: int, pos: PositionDict) -> None:
        """Emit one diagram into *container* (the document or a table cell)."""
        nonlocal bookmark_id
        fen = pos.get('fen', '').strip()
        num = numbers[pos_idx]
        title = _make_title(pos, num, title_template, title_mode, title_custom, opts)
        moves = _clean_moves(pos.get('moves', ''))
        anchor = _anchor_id(pos_idx)
        diag_anchor = f'diag_{anchor}'
        ans_anchor = f'ans_{anchor}'
        side = parse_fen(fen)[1]
        side_indicator = (
            _side_indicator_text(side, symbol)
            if not _has_embedded_side_indicator(font_name) else ''
        )
        diag_lines = fen_to_diagram(
            fen,
            coords=coords,
            flip=flip,
            flip_auto=flip_auto,
            symbol=symbol,
            font_name=font_name,
            border_style=border_style,
        )
        for line in diag_lines:
            used_board_chars.update(line)

        title_pt = 14 if columns == 1 else max(8, 14 - columns)

        if title:
            p_title = container.add_paragraph()
            p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
            _docx_add_bookmark(p_title, diag_anchor, bookmark_id)
            bookmark_id += 1
            if answers_section and moves:
                prefix, suffix = _split_title_number_prefix(title, num, opts)
                if prefix:
                    _docx_add_hyperlink(p_title, prefix, anchor=ans_anchor, bold=True,
                                        font_size_pt=title_pt)
                    if suffix:
                        run = p_title.add_run(suffix)
                        run.bold = True
                        run.font.size = Pt(title_pt)
                else:
                    _docx_add_hyperlink(p_title, title, anchor=ans_anchor, bold=True,
                                        font_size_pt=title_pt)
            else:
                run = p_title.add_run(title)
                run.bold = True
                run.font.size = Pt(title_pt)

        p_board = container.add_paragraph()
        if not title:
            _docx_add_bookmark(p_board, diag_anchor, bookmark_id)
            bookmark_id += 1
        _docx_add_board_paragraph(p_board, diag_lines, font_family_name, board_font_size_pt)

        if side_indicator:
            p_side = container.add_paragraph()
            p_side.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_side.paragraph_format.space_before = Pt(0)
            p_side.paragraph_format.space_after = Pt(6)
            p_side.add_run(side_indicator)

        for line_index in range(lines_count):
            _docx_notation_paragraph(container, line_index, lines_numbered, side, cell_twips - 400)

        if lichess_link:
            p_link = container.add_paragraph()
            p_link.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_link.paragraph_format.space_before = Pt(0)
            _docx_add_hyperlink(p_link, _lichess_label(opts), url=_lichess_analysis_url(fen))

        if not answers_section:
            if moves:
                p_moves = container.add_paragraph()
                p_moves.alignment = WD_ALIGN_PARAGRAPH.LEFT
                p_moves.paragraph_format.space_after = Pt(14)
                nrun = p_moves.add_run(f'{_format_number(num, opts)}. ')
                nrun.bold = True
                _docx_add_moves(p_moves, moves, opts, figurine_family, used_figurine_chars)
        elif moves:
            answers.append((num, moves, diag_anchor, ans_anchor))

    def flush_answers() -> None:
        """Write out the solutions gathered so far, then forget them."""
        nonlocal bookmark_id
        if not answers:
            return
        if opts.answers_new_page:
            doc.add_page_break()
        p_head = doc.add_paragraph()
        p_head.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r_head = p_head.add_run(answers_title)
        r_head.bold = True
        r_head.font.size = Pt(15)

        if answers_cols > 1:
            from docx.enum.section import WD_SECTION
            _docx_set_columns(doc.add_section(WD_SECTION.CONTINUOUS), answers_cols)

        for num, moves, diag_anchor, ans_anchor in answers:
            p_ans = doc.add_paragraph()
            _docx_add_bookmark(p_ans, ans_anchor, bookmark_id)
            bookmark_id += 1
            _docx_add_hyperlink(p_ans, f'{_format_number(num, opts)}. ', anchor=diag_anchor,
                                bold=True, font_size_pt=11)
            _docx_add_moves(p_ans, moves, opts, figurine_family, used_figurine_chars)
        answers.clear()

    for group_index, (chapter, indices) in enumerate(_chapter_groups(valid)):
        if chapter:
            if group_index:
                doc.add_page_break()
            p_chapter = doc.add_paragraph()
            p_chapter.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p_chapter.add_run(chapter)
            run.bold = True
            run.font.size = Pt(16)

        if columns > 1:
            table = doc.add_table(rows=0, cols=columns)
            table.autofit = False
            for row_start in range(0, len(indices), columns):
                row = table.add_row()
                for offset in range(columns):
                    cell = row.cells[offset]
                    # A fresh cell already holds one empty paragraph.
                    cell.paragraphs[0]._p.getparent().remove(cell.paragraphs[0]._p)
                    if row_start + offset >= len(indices):
                        cell.add_paragraph()
                        continue
                    pos_idx = indices[row_start + offset]
                    rendered += 1
                    if progress and total > 0:
                        progress(int(rendered * 100 / total))
                    render_position(cell, pos_idx, valid[pos_idx])
        else:
            for position_number, pos_idx in enumerate(indices):
                rendered += 1
                if progress and total > 0:
                    progress(int(rendered * 100 / total))
                render_position(doc, pos_idx, valid[pos_idx])
                if position_number < len(indices) - 1:
                    doc.add_paragraph()

        if answers_section and opts.answers_after_chapter:
            flush_answers()

    if answers_section:
        flush_answers()

    out_path = Path(out_path)
    doc.save(str(out_path))

    allow_restricted = opts.allow_restricted_fonts
    to_embed = [(font_family_name, docx_font_path, used_board_chars)]
    figurine_path = _figurine_font_path(opts) if figurine_family and used_figurine_chars else None
    if figurine_path is not None:
        to_embed.append((figurine_family, figurine_path, used_figurine_chars))

    for family, path, chars in to_embed:
        if not path.exists() or not check_embedding_permitted(path, 'DOCX', allow_restricted):
            continue
        try:
            _docx_embed_font(out_path, family, path, chars)
        except Exception as exc:
            # The document is already on disk and remains usable through font
            # substitution, so a failure to embed must not look like a failure
            # to generate.
            _warn(f'could not embed {path.name} into the DOCX: {exc}')
